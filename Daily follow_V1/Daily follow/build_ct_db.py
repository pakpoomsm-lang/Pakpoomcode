"""Build CT.db (SQLite) from CT.xlsx.

CT.xlsx (sheet "Master") is the cycle-time master: one row per Material /
Sub Type with the machine assigned to each process and its cycle time (CT).
This script reads that sheet, sorts it into a tidy order, and writes it to a
single `cycle_time` table. Re-run it whenever CT.xlsx changes.

    python build_ct_db.py            # CT.xlsx -> CT.db
"""

import os
import sqlite3
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "CT.xlsx")
DB = os.path.join(HERE, "CT.db")
SHEET = "Master"

# Spreadsheet header -> (column name, SQLite type). Order matches the sheet.
COLUMNS = [
    ("Material", "material", "TEXT"),
    ("Type Code", "type_code", "TEXT"),
    ("Sub Type", "sub_type", "TEXT"),
    ("Sub Type Description", "sub_type_desc", "TEXT"),
    ("Type by sub type", "type_by_sub_type", "TEXT"),
    ("Finpress", "finpress", "TEXT"),
    ("Expander", "expander", "TEXT"),
    ("HPBender", "hpbender", "TEXT"),
    ("Auto-Brazing", "auto_brazing", "TEXT"),
    ("FinPress-Stroke", "finpress_stroke", "REAL"),
    ("Finpress-CT", "finpress_ct", "REAL"),
    ("HairPin-Pcs", "hairpin_pcs", "REAL"),
    ("HPBender-CT", "hpbender_ct", "REAL"),
    ("Expander-Pcs/Cycle", "expander_pcs_per_cycle", "REAL"),
    ("Expander-Sec/Cycle", "expander_sec_per_cycle", "REAL"),
    ("Expander-CT", "expander_ct", "REAL"),
    ("Auto-Brazing-CT", "auto_brazing_ct", "REAL"),
    ("Cutting-Sec", "cutting_sec", "REAL"),
]


def read_rows():
    """Return (header, data_rows) from the Master sheet."""
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    rows = list(wb[SHEET].iter_rows(values_only=True))
    wb.close()
    return rows[0], rows[1:]


def clean(value, sqltype):
    """Normalise a cell: blanks -> None, Material kept as text."""
    if value is None or value == "":
        return None
    if sqltype == "TEXT":
        return str(value).strip()
    return value


def main():
    header, data = read_rows()

    # The sheet column order can drift, so map by header name instead of index.
    idx = {name: i for i, name in enumerate(header)}
    missing = [src for src, _, _ in COLUMNS if src not in idx]
    if missing:
        raise SystemExit(f"CT.xlsx is missing expected columns: {missing}")

    records = []
    for src_row, raw in enumerate(data, start=2):  # row 1 is the header
        rec = [src_row]
        for src, _name, sqltype in COLUMNS:
            rec.append(clean(raw[idx[src]], sqltype))
        records.append(rec)

    # Organise: group by sub type, base rows (no material) first, then by material.
    # Record layout is [src_row, <COLUMNS...>], so a column's index is 1 + its
    # position in COLUMNS.
    names = [c[1] for c in COLUMNS]
    sub_i = 1 + names.index("sub_type")
    mat_i = 1 + names.index("material")
    records.sort(key=lambda r: ((r[sub_i] or ""), (r[mat_i] or "")))

    if os.path.exists(DB):
        os.remove(DB)
    con = sqlite3.connect(DB)
    cols_sql = ",\n  ".join(f"{name} {sqltype}" for _src, name, sqltype in COLUMNS)
    con.execute(f"""
        CREATE TABLE cycle_time (
          src_row INTEGER,
          {cols_sql}
        )
    """)
    placeholders = ",".join("?" * (1 + len(COLUMNS)))
    con.executemany(f"INSERT INTO cycle_time VALUES ({placeholders})", records)
    con.execute("CREATE INDEX idx_ct_material ON cycle_time(material)")
    con.execute("CREATE INDEX idx_ct_sub_type ON cycle_time(sub_type)")
    con.commit()
    con.close()

    print(f"Wrote {len(records)} rows to {DB}")


if __name__ == "__main__":
    main()
