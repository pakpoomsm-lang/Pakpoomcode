"""Build MasterTS.db (SQLite) from MasterTS.xlsx.

MasterTS.xlsx (sheet "Master_TS") is the standard-time master: the TS value
per ITEM that the daily follow uses. This script copies the whole sheet into a
single `master_ts` table, keeping every column. Re-run it whenever
MasterTS.xlsx changes.

    python build_masterts_db.py        # MasterTS.xlsx -> MasterTS.db
"""

import os
import re
import sqlite3
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "MasterTS.xlsx")
DB = os.path.join(HERE, "MasterTS.db")
SHEET = "Master_TS"


def col_name(header):
    """SHOP -> shop, PROCESS_SEQ -> process_seq, 'Foo Bar' -> foo_bar."""
    name = re.sub(r"\W+", "_", str(header).strip().lower()).strip("_")
    return name or "col"


def infer_type(values):
    """REAL if any float, INTEGER if all ints, else TEXT."""
    seen = [v for v in values if v not in (None, "")]
    if not seen:
        return "TEXT"
    if any(isinstance(v, float) for v in seen):
        return "REAL"
    if all(isinstance(v, int) for v in seen):
        return "INTEGER"
    return "TEXT"


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    rows = list(wb[SHEET].iter_rows(values_only=True))
    wb.close()

    header = rows[0]
    data = rows[1:]
    names = [col_name(h) for h in header]
    types = [infer_type([r[i] for r in data]) for i in range(len(header))]

    if os.path.exists(DB):
        os.remove(DB)
    con = sqlite3.connect(DB)
    cols_sql = ",\n  ".join(f"{n} {t}" for n, t in zip(names, types))
    con.execute(f"CREATE TABLE master_ts (\n  src_row INTEGER,\n  {cols_sql}\n)")
    placeholders = ",".join("?" * (1 + len(names)))
    records = [[i] + list(r) for i, r in enumerate(data, start=2)]  # row 1 is header
    con.executemany(f"INSERT INTO master_ts VALUES ({placeholders})", records)
    if "item" in names:
        con.execute("CREATE INDEX idx_master_ts_item ON master_ts(item)")
    con.commit()
    con.close()

    print(f"Wrote {len(records)} rows to {DB}")


if __name__ == "__main__":
    main()
